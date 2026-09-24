from io import BytesIO
import json
from pathlib import Path
import stat
import zipfile

import pytest
from flask import Flask

from app.services.investor_sold_bundle import extract_investor_sold_bundle
from app.services import resumable_uploads
from app.api import investor_sold as api
from app.services import report_store
from openpyxl import load_workbook

BASE = {
    "sold_properties_full.csv": b"sold\n",
    "reisift_export.csv": b"reisift\n",
    "qualified_leads.csv": b"ql\n",
    "sold_property_details.jsonl": b'{"id": 1}\n',
}


def make_zip(tmp_path, files=None, prefix=""):
    path = tmp_path / "bundle.zip"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in (BASE if files is None else files).items():
            archive.writestr(prefix + name, content)
    return path


@pytest.mark.parametrize("prefix", ["", "report/"])
def test_bundle_extracts_streamed_required_and_optional(tmp_path, prefix):
    path = make_zip(tmp_path, {**BASE, "opportunities.xlsx": b"opps", "transactions.csv": b"txns"}, prefix)
    result = extract_investor_sold_bundle(str(path), tmp_path / "extracted")
    assert len(result) == 6
    assert Path(result["property_details_path"]).read_bytes() == BASE["sold_property_details.jsonl"]
    assert Path(result["sold_path"]).parent == tmp_path / "extracted"


def test_bundle_ignores_mac_metadata(tmp_path):
    path = make_zip(tmp_path, {**BASE, "__MACOSX/._whatever": b"metadata", ".DS_Store": b"metadata"})
    assert len(extract_investor_sold_bundle(str(path), tmp_path / "out")) == 4


@pytest.mark.parametrize("extra", ["../escape.csv", "/absolute.csv", "C:/absolute.csv", "folder\\evil.csv", "x/y/qualified_leads.csv", "notes.txt", "qualified_leads.xlsx", "other/qualified_leads.csv"])
def test_bundle_rejects_unsafe_unexpected_or_ambiguous_files(tmp_path, extra):
    path = make_zip(tmp_path, {**BASE, extra: b"x"})
    with pytest.raises(ValueError):
        extract_investor_sold_bundle(str(path), tmp_path / "out")
    assert not (tmp_path / "out").exists()


def test_bundle_rejects_mixed_root_and_wrapper(tmp_path):
    files = dict(BASE)
    files["wrapper/reisift_export.csv"] = files.pop("reisift_export.csv")
    with pytest.raises(ValueError, match="shared wrapping"):
        extract_investor_sold_bundle(str(make_zip(tmp_path, files)), tmp_path / "out")


def test_bundle_rejects_missing_required_sidecar(tmp_path):
    files = {k: v for k, v in BASE.items() if not k.endswith("jsonl")}
    with pytest.raises(ValueError, match="property_details_path"):
        extract_investor_sold_bundle(str(make_zip(tmp_path, files)), tmp_path / "out")


def test_bundle_rejects_symlink(tmp_path):
    path = make_zip(tmp_path)
    with zipfile.ZipFile(path, "a") as archive:
        entry = zipfile.ZipInfo("link")
        entry.create_system = 3
        entry.external_attr = (stat.S_IFLNK | 0o777) << 16
        archive.writestr(entry, "sold_properties_full.csv")
    with pytest.raises(ValueError, match="Unsafe"):
        extract_investor_sold_bundle(str(path), tmp_path / "out")


def test_bundle_limit_is_configurable(tmp_path, monkeypatch):
    monkeypatch.setenv("INVESTOR_SOLD_BUNDLE_MAX_UNCOMPRESSED_BYTES", "5")
    with pytest.raises(ValueError, match="byte limit"):
        extract_investor_sold_bundle(str(make_zip(tmp_path)), tmp_path / "out")


def test_corrupt_bundle_fails_clearly(tmp_path):
    path = tmp_path / "broken.zip"
    path.write_bytes(b"not a zip")
    with pytest.raises(ValueError, match="Invalid or unreadable"):
        extract_investor_sold_bundle(str(path), tmp_path / "out")


def test_crc_failure_removes_already_extracted_files(tmp_path):
    path = tmp_path / "bad-crc.zip"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_STORED) as archive:
        for name, content in BASE.items():
            archive.writestr(name, content)
    data = path.read_bytes()
    assert BASE["sold_property_details.jsonl"] in data
    path.write_bytes(data.replace(BASE["sold_property_details.jsonl"], b'{"id": 2}\n', 1))
    with pytest.raises(ValueError, match="CRC"):
        extract_investor_sold_bundle(str(path), tmp_path / "out")
    assert list((tmp_path / "out").iterdir()) == []


def test_duplicate_archive_entry_rejected(tmp_path):
    path = make_zip(tmp_path)
    with zipfile.ZipFile(path, "a") as archive:
        with pytest.warns(UserWarning, match="Duplicate name"):
            archive.writestr("sold_properties_full.csv", "duplicate")
    with pytest.raises(ValueError, match="Ambiguous"):
        extract_investor_sold_bundle(str(path), tmp_path / "out")


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(api, "IS_ROOT", tmp_path / "jobs")
    monkeypatch.setattr(resumable_uploads, "_FINAL_DIR", tmp_path)
    app = Flask(__name__)
    app.register_blueprint(api.investor_sold_bp, url_prefix="/gate7")
    return app.test_client()


def capture_start(monkeypatch):
    calls = []
    def start(*args, **kwargs):
        calls.append((args, kwargs))
        return {"job_id": "test", "status": "started"}, 202
    monkeypatch.setattr(api, "_start_investor_sold_job", start)
    return calls


def test_bundle_json_queues_without_extracting(client, tmp_path, monkeypatch):
    # Invalid bytes are deliberately accepted at request time; worker validates ZIP.
    path = tmp_path / "bundle.zip"
    path.write_bytes(b"worker validates")
    calls = capture_start(monkeypatch)
    response = client.post("/gate7/analyze", json={"bundle_path": str(path)})
    assert response.status_code == 202
    assert calls[0][1]["bundle_path"] == str(path.resolve())


def test_bundle_json_requires_trusted_path(client, tmp_path, monkeypatch):
    calls = capture_start(monkeypatch)
    response = client.post("/gate7/analyze", json={"bundle_path": str(tmp_path.parent / "outside.zip")})
    assert response.status_code == 400
    assert "completed resumable" in response.json["detail"]
    assert not calls


def test_bundle_cannot_mix_individual_paths(client, monkeypatch):
    calls = capture_start(monkeypatch)
    response = client.post("/gate7/analyze", json={"bundle_path": "a.zip", "sold_path": "b.csv"})
    assert response.status_code == 400
    assert not calls


def test_individual_details_path_is_resolved_and_forwarded(client, tmp_path, monkeypatch):
    calls = capture_start(monkeypatch)
    body = {}
    for role, filename in [("sold_path", "sold.csv"), ("reisift_path", "rei.csv"), ("ql_path", "ql.csv"), ("property_details_path", "details.jsonl")]:
        path = tmp_path / filename
        path.write_text("data")
        body[role] = str(path)
    assert client.post("/gate7/analyze", json=body).status_code == 202
    assert calls[0][1]["property_details_path"] == str((tmp_path / "details.jsonl").resolve())
    body["property_details_path"] = str(tmp_path.parent / "outside.jsonl")
    assert client.post("/gate7/analyze", json=body).status_code == 400


def test_multipart_bundle_alone(client, monkeypatch):
    calls = capture_start(monkeypatch)
    response = client.post("/gate7/analyze", data={"bundle_file": (BytesIO(b"zip"), "inputs.zip")})
    assert response.status_code == 202
    assert Path(calls[0][1]["bundle_path"]).read_bytes() == b"zip"


def test_multipart_individual_same_names_do_not_overwrite(client, monkeypatch):
    calls = capture_start(monkeypatch)
    response = client.post("/gate7/analyze", data={
        "sold_file": (BytesIO(b"sold"), "same.csv"),
        "reisift_file": (BytesIO(b"rei"), "same.csv"),
        "ql_file": (BytesIO(b"ql"), "same.csv"),
        "property_details_file": (BytesIO(b"details"), "details.jsonl"),
    })
    assert response.status_code == 202
    args, kwargs = calls[0]
    assert [Path(p).read_bytes() for p in args[:3]] == [b"sold", b"rei", b"ql"]
    assert Path(kwargs["property_details_path"]).read_bytes() == b"details"


@pytest.mark.parametrize("kind,filename", [("investor_sold_bundle", "a.zip"), ("property_details", "a.jsonl")])
def test_resumable_new_kinds(kind, filename):
    assert resumable_uploads._validate_kind(kind) == kind
    assert resumable_uploads._validate_filename(kind, filename) == filename
    with pytest.raises(ValueError):
        resumable_uploads._validate_filename(kind, "bad.csv")


def test_worker_extracts_and_passes_details(tmp_path, monkeypatch):
    received = {}
    class Result:
        def to_api_dict(self):
            return {"warnings": []}
    def analyze(sold_path, **kwargs):
        received.update(kwargs)
        assert Path(sold_path).read_bytes() == BASE["sold_properties_full.csv"]
        return Result()
    monkeypatch.setattr(api, "analyze", analyze)
    monkeypatch.setattr(api, "save_investor_sold_report", lambda *args, **kwargs: None)
    job = tmp_path / "job"
    api._analyze_in_subprocess("test", "", None, None, None, None, str(job), bundle_path=str(make_zip(tmp_path)))
    assert Path(received["property_details_path"]).read_bytes() == BASE["sold_property_details.jsonl"]
    assert api._read_job_progress(job)["status"] == "completed"


def test_invalid_bundle_marks_job_failed_without_analyzing(tmp_path, monkeypatch):
    path = tmp_path / "bad.zip"
    path.write_bytes(b"broken")
    monkeypatch.setattr(api, "analyze", lambda *a, **kw: pytest.fail("invalid bundle reached analyze"))
    job = tmp_path / "job"
    api._analyze_in_subprocess("test", "", None, None, None, None, str(job), bundle_path=str(path))
    progress = api._read_job_progress(job)
    assert progress["status"] == "failed"
    assert "ZIP bundle" in progress["message"]


def test_job_launch_preserves_positional_job_id_and_appends_inputs(tmp_path, monkeypatch):
    captured = {}
    class Thread:
        def __init__(self, **kwargs):
            captured.update(kwargs)
        def start(self):
            pass
    monkeypatch.setattr(api, "IS_ROOT", tmp_path)
    monkeypatch.setattr(api.threading, "Thread", Thread)
    payload, status = api._start_investor_sold_job("sold", "rei", "ql", None, "txn", "existing-job", property_details_path="details", bundle_path=None)
    assert status == 202
    assert payload["job_id"] == "existing-job"
    assert captured["args"] == ("existing-job", "sold", "rei", "ql", None, "txn", "details", None)
    api._jobs.pop("existing-job", None)


def test_real_zip_worker_analysis_saved_reload_and_export(tmp_path, monkeypatch):
    fixtures = Path(__file__).parent / "fixtures"
    sold = (
        "period_date,period_label,dataflik_id,transaction_id,parcel_number,buyer_full_name,"
        "property_address,property_city,property_zip,county,state,sale_amount,investor,in_my_records,investor_score\n"
        "2025-03-01,Mar 2025,1,t1,12-34,Example Buyer LLC,100 Main St,Hempstead,11550,Nassau,NY,500000,TRUE,FALSE,80\n"
        "2025-04-01,Apr 2025,1,t2,12-34,Later Buyer LLC,100 Main St,Hempstead,11550,Nassau,NY,600000,FALSE,TRUE,20\n"
    )
    details = {
        "dataflik_id": "1",
        "property": {"apn": "1234", "county": "Nassau County", "state": "NY",
                     "property_type": "single_family_residence", "years_built": 2000,
                     "total_market_value": 500000, "living_square_feet": 1000, "lot_sqrf": 4000},
        "raw": {"detail": {"units_count": 1}},
        "history": {"transactions": [{"sale_date": "2025-03-15", "buyer_name": "Example Buyer LLC",
                                      "sale_price": 500000, "seller_name": "Smith Family Trust"}]},
    }
    bundle = make_zip(tmp_path, {
        "sold_properties_full.csv": sold.encode(),
        "reisift_export.csv": (fixtures / "investor_sold_reisift.csv").read_bytes(),
        "qualified_leads.csv": (fixtures / "investor_sold_ql.csv").read_bytes(),
        "sold_property_details.jsonl": (json.dumps(details) + "\n").encode(),
    }, prefix="report/")
    saved_root = tmp_path / "saved_reports"
    def save(job_id, metrics, created_at):
        return report_store.save_investor_sold_report(job_id, metrics, created_at, reports_dir=saved_root)
    monkeypatch.setattr(api, "save_investor_sold_report", save)
    job = tmp_path / "job"
    api._analyze_in_subprocess("real-test", "", None, None, None, None, str(job), bundle_path=str(bundle))
    progress = api._read_job_progress(job)
    assert progress["status"] == "completed", progress
    persisted = report_store.load_investor_sold_report("real-test", reports_dir=saved_root)
    metrics = persisted["metrics"]
    assert metrics["inputs"]["property_rows"] == 1
    assert metrics["property_screening"]["eligible"] == 1
    assert metrics["property_screening"]["seller_categories"] == {"Trust": 1}
    audit = metrics["property_screening_rows"][0]
    assert audit["seller_name"] == "Smith Family Trust"
    assert audit["seller_match_status"] == "matched_sale_history"
    restored = api.result_from_metrics_dict(metrics)
    assert restored.rows[0].seller_category == "Trust"
    assert restored.rows[0].buyer_full_name == "Example Buyer LLC"
    assert restored.rows[0].transaction_count == 2
    workbook = load_workbook(BytesIO(api.build_export_workbook(restored)), read_only=True)
    try:
        exported_values = [str(value) for sheet in workbook for row in sheet.iter_rows(values_only=True) for value in row]
        assert "Smith Family Trust" in exported_values
        assert "Trust" in exported_values
    finally:
        workbook.close()
